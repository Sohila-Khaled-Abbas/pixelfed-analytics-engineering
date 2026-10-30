"""Read native exports and verify keys after serialization, without repairing IDs."""

from __future__ import annotations

import csv
import gzip
import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def parse_timestamp(value):
    """Only documented ISO-8601 and controlled Bronze slash format; naive means UTC."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value).strip()
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            parsed = datetime.strptime(text, "%Y/%m/%d %H:%M:%S")
    return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed.astimezone(timezone.utc)


def read_tables(landing: Path, storage: dict) -> dict[str, list[dict[str, Any]]]:
    """Read the manifest's six native formats; reject paths outside landing."""
    data = {}
    for table, info in storage.items():
        path = (landing / info["file"]).resolve()
        if not path.is_relative_to(landing.resolve()):
            raise ValueError(f"{table}: path escapes landing")
        fmt = info["format"]
        if fmt in {"csv", "csv_gz"}:
            opener = gzip.open if fmt == "csv_gz" else open
            with opener(path, "rt", encoding="utf-8", newline="") as stream:
                data[table] = list(csv.DictReader(stream))
        elif fmt == "json":
            data[table] = json.loads(path.read_text(encoding="utf-8"))
        elif fmt == "jsonl":
            data[table] = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
        elif fmt == "xlsx":
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = book[table[:31]].iter_rows(values_only=True)
                header = next(rows)
                data[table] = []
                for row in rows:
                    if len(row) > len(header):
                        raise ValueError(f"{table}: XLSX row exceeds schema width")
                    # XLSX omits trailing blank cells physically; preserve their logical NULLs.
                    padded = (*row, *(None for _ in range(len(header) - len(row))))
                    data[table].append(dict(zip(header, padded, strict=True)))
            finally:
                book.close()
        elif fmt == "sqlite":
            with sqlite3.connect(path.as_uri() + "?mode=ro", uri=True) as conn:
                conn.row_factory = sqlite3.Row
                quoted = '"' + table.replace('"', '""') + '"'
                data[table] = [dict(row) for row in conn.execute(f"SELECT * FROM {quoted}")]
        else:
            raise ValueError(f"Unsupported format: {fmt}")
    return data


def key(value: Any) -> str | None:
    """Normalize exported NULLs and integers exactly; never round through float."""
    if value is None or value == "":
        return None
    if isinstance(value, float):
        # Excel/float cannot preserve larger integers exactly.
        if not value.is_integer() or abs(value) > 2**53 - 1:
            raise ValueError(f"Unsafe floating-point identifier: {value}")
        return str(int(value))
    return str(value)


def validate_records(data: dict, schema: Any) -> list[str]:
    """Recheck PK/unique/FK and canonical identity across physical formats."""
    errors = []
    normalized = {}
    for name, spec in schema.tables.items():
        normalized[name] = []
        for row in data.get(name, []):
            if list(row) != spec.column_names:
                errors.append(f"{name}: record columns/order mismatch")
            clean = {}
            for col in spec.columns:
                value = row.get(col.name)
                integer = bool(re.match(r"^(?:tinyint|smallint|mediumint|bigint|int)\b", col.mysql_type))
                value = (
                    key(value) if integer else (None if value is None or (value == "" and col.nullable) else str(value))
                )
                clean[col.name] = value
                if value is None and (col.primary_key or not col.nullable):
                    errors.append(f"{name}.{col.name}: required value missing")
                if value is not None and re.match(r"^(?:tinyint|smallint|mediumint|bigint|int)\b", col.mysql_type):
                    if not re.fullmatch(r"-?\d+", value):
                        errors.append(f"{name}.{col.name}: invalid integer {value!r}")
            normalized[name].append(clean)
        for col in spec.columns:
            if col.primary_key or col.unique:
                values = [row[col.name] for row in normalized[name] if row[col.name] is not None]
                if len(values) != len(set(values)):
                    errors.append(f"{name}.{col.name}: duplicate key")
    for table, columns, unique, _ in schema.indexes:
        if unique:
            values = [tuple(row[c] for c in columns) for row in normalized[table]]
            values = [v for v in values if None not in v]
            if len(values) != len(set(values)):
                errors.append(f"{table}: duplicate unique tuple {columns}")
    # Apply the same documented source-schema relationship correction as generation.
    from schema_generator_core import build_reference_map

    for (child, col), (parent, target) in build_reference_map(schema).items():
        allowed = {row[target] for row in normalized[parent]}
        if any(row[col] is not None and row[col] not in allowed for row in normalized[child]):
            errors.append(f"{child}.{col}: orphan referencing {parent}.{target}")
    users = {r["id"]: r for r in normalized.get("users", [])}
    profiles = {r["id"]: r for r in normalized.get("profiles", [])}
    posts = {r["id"]: r for r in normalized.get("statuses", [])}
    for uid, user in users.items():
        profile = profiles.get(user["profile_id"])
        if profile is None or profile["user_id"] != uid:
            errors.append(f"users.{uid}: bidirectional identity mismatch")
        elif (user["name"], user["username"]) != (profile["name"], profile["username"]):
            errors.append(f"users.{uid}: persona mismatch")
    for pid, profile in profiles.items():
        uid = profile["user_id"]
        if uid is not None and users.get(uid, {}).get("profile_id") != pid:
            errors.append(f"profiles.{pid}: bidirectional identity mismatch")
        if (uid is None) != (profile["domain"] is not None):
            errors.append(f"profiles.{pid}: local/federated identity mismatch")
    for table in ("comments", "media", "reports", "report_comments"):
        for row in normalized.get(table, []):
            profile = profiles.get(row["profile_id"])
            if profile and row["user_id"] != profile["user_id"]:
                errors.append(f"{table}.{row['id']}: actor identity mismatch")
    for table in ("media", "status_hashtags"):
        for row in normalized.get(table, []):
            post = posts.get(row["status_id"])
            if post and post["profile_id"] is not None and row["profile_id"] != post["profile_id"]:
                errors.append(f"{table}.{row['id']}: post author mismatch")
    return errors


def validate_event_chronology(data: dict, as_of=None) -> list[str]:
    """Events follow actor registration and target creation; story views precede expiry."""
    errors = []
    profiles = {key(r["id"]): r for r in data.get("profiles", [])}
    posts = {key(r["id"]): r for r in data.get("statuses", [])}
    stories = {key(r["id"]): r for r in data.get("stories", [])}
    cutoff = parse_timestamp(as_of) if as_of is not None else None
    for table in ("statuses", "likes", "comments", "bookmarks", "followers", "stories", "story_views", "media"):
        for row in data.get(table, []):
            try:
                event = parse_timestamp(row.get("created_at"))
                if event is None:
                    continue  # Nullable source time; analytics builder quarantines undated facts.
                parents = [profiles.get(key(row.get("profile_id")), {})]
                if table == "followers":
                    parents.append(profiles.get(key(row.get("following_id")), {}))
                if table in {"likes", "comments", "bookmarks", "media"}:
                    parents.append(posts.get(key(row.get("status_id")), {}))
                if table == "story_views":
                    story = stories.get(key(row.get("story_id")), {})
                    parents.append(story)
                    expiry = parse_timestamp(story.get("expires_at"))
                    if expiry is not None and event > expiry:
                        errors.append(f"{table}.{row['id']}: event after story expiry")
                if any(
                    parse_timestamp(p.get("created_at")) is not None and event < parse_timestamp(p["created_at"])
                    for p in parents
                ):
                    errors.append(f"{table}.{row['id']}: event before actor or target creation")
                if cutoff is not None and event > cutoff:
                    errors.append(f"{table}.{row['id']}: event after cutoff")
            except (ValueError, TypeError):
                errors.append(f"{table}.{row.get('id')}: invalid event timestamp")
    return errors
